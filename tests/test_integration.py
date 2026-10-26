"""
End-to-End Integration Tests.
Verifies complete flow from ingestion -> cleaning -> aggregation -> DB -> forecasting -> inventory -> Excel report.
"""

import pytest
import io
import pandas as pd
import openpyxl

from src.ingestion.column_mapper import auto_detect_columns
from src.preprocessing.cleaning import DataCleaner
from src.preprocessing.aggregation import aggregate_daily_sales
from src.forecasting.train import DemandModelTrainer
from src.forecasting.predict import generate_future_forecast
from src.inventory.recommendations import generate_mode_b_recommendations
from src.exports.excel_report import generate_retail_master_excel
from src.analytics.kpis import compute_overview_kpis


def test_full_pipeline_integration(in_memory_repo):
    repo = in_memory_repo
    shop = repo.get_shop(1)

    # 1. Ingestion: Raw transactions with messy column headers
    raw_csv = """Transaction Date,Item SKU,Item Description,Dept,Qty Sold,Price Each,Total Amount
2026-09-01,SKU-A,Basmati Rice,Groceries,4,100,400
2026-09-01,SKU-B,Sunflower Oil,Groceries,2,150,300
2026-09-02,SKU-A,Basmati Rice,Groceries,5,100,500
2026-09-03,SKU-A,Basmati Rice,Groceries,-1,100,-100
2026-09-04,SKU-A,Basmati Rice,Groceries,6,100,600
2026-09-05,SKU-A,Basmati Rice,Groceries,7,100,700
2026-09-06,SKU-A,Basmati Rice,Groceries,8,100,800
2026-09-07,SKU-A,Basmati Rice,Groceries,9,100,900
2026-09-08,SKU-A,Basmati Rice,Groceries,10,100,1000
2026-09-09,SKU-A,Basmati Rice,Groceries,8,100,800
2026-09-10,SKU-A,Basmati Rice,Groceries,7,100,700
2026-09-11,SKU-A,Basmati Rice,Groceries,6,100,600
2026-09-12,SKU-A,Basmati Rice,Groceries,5,100,500
2026-09-13,SKU-A,Basmati Rice,Groceries,6,100,600
2026-09-14,SKU-A,Basmati Rice,Groceries,7,100,700
2026-09-15,SKU-A,Basmati Rice,Groceries,8,100,800
2026-09-16,SKU-A,Basmati Rice,Groceries,9,100,900
"""
    raw_df = pd.read_csv(io.StringIO(raw_csv))

    # 2. Auto-detect columns
    mapping = auto_detect_columns(list(raw_df.columns))
    assert mapping["date"] == "Transaction Date"
    assert mapping["product_id"] == "Item SKU"
    assert mapping["quantity"] == "Qty Sold"

    # 3. Clean data
    cleaner = DataCleaner(mapping)
    clean_df, quality_report = cleaner.clean_sales_data(raw_df, allow_returns=True)
    assert quality_report["accepted_rows"] == len(clean_df)
    assert quality_report["negative_quantities"] == 1

    # 4. Aggregate to daily
    daily_df = aggregate_daily_sales(clean_df, shop_id=1, fill_gaps_with_zeros=True)
    assert not daily_df.empty

    # 5. Save to database
    repo.upsert_products(1, [{
        "product_id": "SKU-A",
        "product_name": "Basmati Rice",
        "category": "Groceries",
        "unit_price": 100.0,
        "lead_time_days": 3,
        "minimum_order_quantity": 5.0,
        "case_pack_size": 5.0
    }])
    upload_id = repo.create_upload_record(1, "raw.csv", "hash", "sales", len(raw_df), quality_report)
    repo.save_daily_sales(1, upload_id, daily_df)

    # Save inventory snapshot
    repo.save_inventory_snapshot(1, [{
        "product_id": "SKU-A",
        "on_hand_qty": 6.0,
        "on_order_qty": 0.0,
        "backorders": 0.0,
        "lead_time_days": 3
    }])

    # 6. Forecasting with chronological validation
    trainer = DemandModelTrainer(val_days=4)
    sku_a_daily = daily_df[daily_df["product_id"] == "SKU-A"].sort_values(by="sale_date").reset_index(drop=True)
    res = trainer.train_and_evaluate(sku_a_daily, model_preference="Best Model (Auto-select)")
    assert "val_metrics" in res

    forecast_df = generate_future_forecast(res, product_id="SKU-A", horizon_days=7)
    assert len(forecast_df) == 7

    run_id = repo.save_forecast_run(
        shop_id=1,
        model_name=res["selected_model_name"],
        horizon_days=7,
        validation_metrics=res["val_metrics"],
        config={},
        forecast_rows=forecast_df.to_dict(orient="records")
    )

    # 7. Inventory Optimization (Mode B)
    snaps_df = repo.get_latest_inventory_snapshots(1)
    recs = generate_mode_b_recommendations(
        forecast_results_df=forecast_df,
        inventory_snapshots_df=snaps_df,
        historical_daily_sales_df=sku_a_daily,
        review_period_days=7,
        safety_stock_method="buffer",
        buffer_percentage=0.20
    )
    assert len(recs) == 1
    assert recs[0]["suggested_qty"] > 0
    repo.save_recommendations(run_id, recs)

    # 8. Master Excel Export
    recs_df = repo.get_recommendations_df(run_id)
    kpis = compute_overview_kpis(daily_df)
    prods_df = repo.get_products_df(1)

    excel_bytes = generate_retail_master_excel(
        shop_info=shop,
        kpis=kpis,
        cleaned_sales_df=daily_df,
        quality_report=quality_report,
        product_analytics_df=prods_df,
        forecast_results_df=forecast_df,
        validation_metrics=res["val_metrics"],
        recommendations_df=recs_df,
        inventory_mode="inventory_mode"
    )

    assert len(excel_bytes) > 1000
    # Verify openpyxl can read it and all sheets are present
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), read_only=True)
    assert "Executive Summary" in wb.sheetnames
    assert "Cleaned Sales" in wb.sheetnames
    assert "Data Quality Audit" in wb.sheetnames
    assert "Demand Forecast" in wb.sheetnames
    assert "Inventory Recommendations" in wb.sheetnames
    wb.close()
