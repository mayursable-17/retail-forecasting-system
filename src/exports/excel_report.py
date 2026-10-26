"""
Professional Excel & CSV Report Generation Engine.
Creates multi-tab Excel workbooks with custom styles, numbers, and executive summaries.
"""

import io
from datetime import datetime, timezone
from typing import Dict, Any, Optional
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def format_excel_sheet(ws, title_header: str = None):
    """Apply styling, header colors, borders, and auto-adjust column widths."""
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    border_thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    # Style header row (assumed row 1)
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Borders and formatting for data rows
    for row_idx in range(2, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border_thin
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="right")
            elif isinstance(cell.value, (datetime, pd.Timestamp)):
                cell.number_format = "YYYY-MM-DD"
                cell.alignment = Alignment(horizontal="center")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(12, min(max_len + 4, 50))


def generate_retail_master_excel(
    shop_info: Dict[str, Any],
    kpis: Dict[str, Any],
    cleaned_sales_df: pd.DataFrame,
    quality_report: Optional[Dict[str, Any]] = None,
    product_analytics_df: Optional[pd.DataFrame] = None,
    forecast_results_df: Optional[pd.DataFrame] = None,
    validation_metrics: Optional[Dict[str, Any]] = None,
    recommendations_df: Optional[pd.DataFrame] = None,
    inventory_mode: str = "inventory_mode"
) -> bytes:
    """
    Generate comprehensive retail report workbook as raw bytes for download.
    """
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        wb = writer.book

        # Tab 1: Executive Summary
        ws_summary = wb.create_sheet(title="Executive Summary", index=0)
        ws_summary.views.sheetView[0].showGridLines = True

        title_font = Font(name="Calibri", size=16, bold=True, color="1F497D")
        sub_font = Font(name="Calibri", size=11, italic=True)
        bold_font = Font(name="Calibri", size=11, bold=True)
        val_font = Font(name="Calibri", size=11)

        ws_summary["A1"] = f"{shop_info.get('shop_name', 'Retail Store')} - System Analytics & Inventory Report"
        ws_summary["A1"].font = title_font
        ws_summary["A2"] = f"Generated on: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')} | Currency: {shop_info.get('currency_code', 'INR')}"
        ws_summary["A2"].font = sub_font

        ws_summary["A4"] = "Core Metric"
        ws_summary["B4"] = "Value"
        ws_summary["A4"].font = bold_font
        ws_summary["B4"].font = bold_font
        ws_summary["A4"].fill = PatternFill(start_color="D9E1F2", fill_type="solid")
        ws_summary["B4"].fill = PatternFill(start_color="D9E1F2", fill_type="solid")

        cur = shop_info.get("currency_symbol", "₹")
        summary_rows = [
            ("Store Type", shop_info.get("shop_type", "Retail")),
            ("Inventory Optimization Mode", "Workflow B (Inventory + Sales)" if inventory_mode == "inventory_mode" else "Workflow A (Sales Only - Estimated)"),
            ("Analysis Date Range", f"{kpis.get('start_date', 'N/A')} to {kpis.get('end_date', 'N/A')} ({kpis.get('total_days', 0)} days)"),
            ("Total Catalog Products", kpis.get("total_products", 0)),
            ("Total Units Sold", f"{kpis.get('total_units_sold', 0.0):,.1f}"),
            ("Total Historical Revenue", f"{cur}{kpis.get('total_revenue', 0.0):,.2f}"),
            ("Average Daily Revenue", f"{cur}{kpis.get('avg_daily_revenue', 0.0):,.2f}"),
        ]

        if validation_metrics:
            summary_rows.extend([
                ("Forecasting Validation Model", validation_metrics.get("model_name", "N/A")),
                ("Mean Absolute Error (MAE)", f"{validation_metrics.get('mae', 0.0):.3f}"),
                ("Weighted Abs % Error (WAPE)", f"{validation_metrics.get('wape_pct', 0.0):.2f}%"),
                ("Forecast Bias", f"{validation_metrics.get('bias_pct', 0.0):.2f}%")
            ])

        for r_idx, (label, val) in enumerate(summary_rows, start=5):
            c1 = ws_summary.cell(row=r_idx, column=1, value=label)
            c2 = ws_summary.cell(row=r_idx, column=2, value=str(val))
            c1.font = bold_font
            c2.font = val_font

        ws_summary.column_dimensions["A"].width = 35
        ws_summary.column_dimensions["B"].width = 45

        # Tab 2: Cleaned Sales Data
        if not cleaned_sales_df.empty:
            cleaned_sales_df.to_excel(writer, sheet_name="Cleaned Sales", index=False)
            format_excel_sheet(writer.sheets["Cleaned Sales"])

        # Tab 3: Data Quality Report
        if quality_report:
            q_rows = [
                {"Audit Parameter": k.replace("_", " ").title(), "Metric / Action": str(v)}
                for k, v in quality_report.items() if k != "cleaning_actions"
            ]
            if "cleaning_actions" in quality_report:
                for act in quality_report["cleaning_actions"]:
                    q_rows.append({"Audit Parameter": "Cleaning Action Applied", "Metric / Action": act})
            q_df = pd.DataFrame(q_rows)
            q_df.to_excel(writer, sheet_name="Data Quality Audit", index=False)
            format_excel_sheet(writer.sheets["Data Quality Audit"])

        # Tab 4: Product Analytics
        if product_analytics_df is not None and not product_analytics_df.empty:
            product_analytics_df.to_excel(writer, sheet_name="Product Analytics", index=False)
            format_excel_sheet(writer.sheets["Product Analytics"])

        # Tab 5: Forecast Results
        if forecast_results_df is not None and not forecast_results_df.empty:
            forecast_results_df.to_excel(writer, sheet_name="Demand Forecast", index=False)
            format_excel_sheet(writer.sheets["Demand Forecast"])

        # Tab 6: Inventory Recommendations
        if recommendations_df is not None and not recommendations_df.empty:
            recommendations_df.to_excel(writer, sheet_name="Inventory Recommendations", index=False)
            format_excel_sheet(writer.sheets["Inventory Recommendations"])

            # Tab 7: Critical Stock Risks
            stock_risks = recommendations_df[recommendations_df["risk_status"].isin(["Stock Risk", "Reorder Review"])]
            if not stock_risks.empty:
                stock_risks.to_excel(writer, sheet_name="Stock Risk Alerts", index=False)
                format_excel_sheet(writer.sheets["Stock Risk Alerts"])

    output.seek(0)
    return output.getvalue()


def convert_df_to_csv_bytes(df: pd.DataFrame) -> bytes:
    """Convert dataframe to utf-8 CSV bytes."""
    return df.to_csv(index=False).encode("utf-8")
